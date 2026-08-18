from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


def extract_text_from_file(
    file_path: str,
    file_type: str,
) -> str:
    """
    Extract readable text from a supported knowledge-base file.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    file_type = file_type.upper()

    # -------------------------
    # PDF
    # -------------------------
    if file_type == "PDF":
        reader = PdfReader(str(path))

        pages = []

        for page in reader.pages:
            text = page.extract_text() or ""

            if text.strip():
                pages.append(text)

        return "\n\n".join(pages)

    # -------------------------
    # DOCX
    # -------------------------
    if file_type == "DOCX":
        document = Document(str(path))

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        return "\n".join(paragraphs)

    # -------------------------
    # XLSX
    # -------------------------
    if file_type == "XLSX":
        workbook = load_workbook(
            filename=str(path),
            read_only=True,
            data_only=True,
        )

        sheets = []

        for worksheet in workbook.worksheets:
            rows = []

            for row in worksheet.iter_rows(
                values_only=True
            ):
                values = []

                for value in row:
                    if value is not None:
                        values.append(str(value))

                if values:
                    rows.append(" | ".join(values))

            if rows:
                sheet_text = (
                    f"Sheet: {worksheet.title}\n"
                    + "\n".join(rows)
                )

                sheets.append(sheet_text)

        workbook.close()

        return "\n\n".join(sheets)

    # -------------------------
    # TXT
    # -------------------------
    if file_type == "TXT":
        return path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

    raise ValueError(
        f"Unsupported file type: {file_type}"
    )